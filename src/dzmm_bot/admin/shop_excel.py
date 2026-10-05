from base64 import b64encode
from io import BytesIO
from math import isfinite
from zipfile import ZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from starlette.exceptions import HTTPException
from starlette.responses import JSONResponse

from dzmm_bot.core.shop_effects import EFFECT_LABELS, PARAMETER_LABELS, effect_dictionary
from dzmm_bot.core.shop_management import FIELD_LABELS, shop_issue


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
COLUMNS = (
    "operation", "public_number", "name", "description", "price", "stock",
    "unlimited_stock", "enabled", "category", "daily_purchase_limit", "minimum_rank_order",
    "effect_type", *PARAMETER_LABELS, "system_key", "id", "configuration_version", "base_stock", "deleted_at",
)
READ_ONLY = {"system_key", "id", "configuration_version", "base_stock", "deleted_at"}
REQUIRED = set(COLUMNS) - READ_ONLY - set(PARAMETER_LABELS)
REPORT_COLUMNS = ("校验状态", "错误原因", "修改建议")
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ShopExcelUploadLimitMiddleware:
    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope["type"] != "http" or scope.get("path") != "/api/game/shop/excel/preview" or scope.get("method") != "POST":
            return await self.app(scope, receive, send)
        maximum = MAX_UPLOAD_BYTES + 64 * 1024
        detail = "上传文件或表单过大；文件不能超过 5MB，请缩小文件并只提交必要字段"
        for name, value in scope.get("headers", []):
            if name == b"content-length":
                try:
                    length = int(value)
                    if length < 0:
                        raise ValueError()
                except ValueError:
                    return await JSONResponse({"detail": "上传请求长度无效，请重新选择文件上传"}, status_code=400)(scope, receive, send)
                if length > maximum:
                    return await JSONResponse({"detail": detail}, status_code=413)(scope, receive, send)
        received = 0

        async def limited_receive():
            nonlocal received
            message = await receive()
            received += len(message.get("body", b""))
            if received > maximum:
                raise HTTPException(413, detail)
            return message

        return await self.app(scope, limited_receive, send)


def _text_cell(cell, value):
    if isinstance(value, str):
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value


def _workbook(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "商品列表"
    worksheet.append([FIELD_LABELS[field] for field in COLUMNS])
    for line, item in enumerate(rows, 2):
        config = item.get("effect_config") or {}
        for column, field in enumerate(COLUMNS, 1):
            value = config.get(field) if field in PARAMETER_LABELS else item.get(field)
            if field == "operation":
                value = "跳过" if item.get("deleted_at") else "更新"
            elif field == "effect_type":
                value = EFFECT_LABELS[value or "none"]
            elif field in ("unlimited_stock", "enabled"):
                value = "是" if value else "否"
            elif field == "base_stock":
                value = item.get("stock")
            _text_cell(worksheet.cell(line, column), value)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column, field in enumerate(COLUMNS, 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 45 if field == "description" else 20
        cell = worksheet.cell(1, column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="666666" if field in READ_ONLY else "254D71")
    for values, field in (("新增,更新,删除,恢复,跳过", "operation"), ("是,否", "enabled"), ("是,否", "unlimited_stock"), (",".join(EFFECT_LABELS.values()), "effect_type")):
        validation = DataValidation(type="list", formula1=f'"{values}"', allow_blank=False)
        validation.errorTitle = "填写值无效"
        validation.error = "请选择列表中的值"
        validation.showErrorMessage = True
        worksheet.add_data_validation(validation)
        letter = get_column_letter(COLUMNS.index(field) + 1)
        validation.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
    notes = workbook.create_sheet("填写说明")
    notes.append(["项目", "规则与修改建议"])
    instructions = (
        ("流程", "先校验和预览，全部错误修正、风险逐项确认后才整批执行。"),
        ("新增", "操作填新增，编号及只读标识留空；名称、描述必填，名称唯一。"),
        ("更新", "按编号更新；可改名，不能修改商品ID、系统标识、配置版本及库存基准。"),
        ("删除", "必须明确填删除。缺少一行不代表删除；有人持有或流程未结束时请先下架。"),
        ("恢复", "按原编号恢复，默认下架；恢复后再更新配置。"),
        ("跳过", "此行不执行操作。已删除商品导出默认跳过，不会因回导而意外恢复。"),
        ("权限", "删除、恢复、修改效果仅超级管理员可操作。"),
        ("金额和库存", "价格 0–999，库存 0–99999，全部填写整数。已有库存默认不覆盖，新增使用填写库存。"),
        ("名称和描述", "名称 1–64 字，描述 1–200 字，分类最多 32 字；不接受公式。"),
        ("每日限购", "0 或留空不限，范围 0–99；同分类共用额度。"),
        ("最低职位", "填当前后台有效职位序号；留空不限制。"),
        ("效果", "从效果字典选择类型和同类型模板。适用参数必须填写，不适用的参数清空。"),
        ("效果参数", "币额 0–99999；次数 1–999；受邀人数 1–5；持续分钟数 1–10080。"),
        ("效果变更", "已有未使用背包采用新效果，进行中的流程保留原快照。"),
        ("错误报告", "修正标出的错误后可直接重新上传；校验状态、错误原因、修改建议列不参与导入。"),
        ("示例新增", "新增 / 编号留空 / 名称：示例赠送卡 / 描述：赠送2币 / 价格：3 / 库存：10 / 无限库存：否 / 上架状态：是 / 效果类型：赠送币 / 赠送金额：2。"),
    )
    for row in instructions:
        notes.append(row)
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 100
    for row in notes:
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    dictionary = workbook.create_sheet("效果字典")
    dictionary.append(["效果类型", "类型代码", "必填参数", "模板代码", "模板名称"])
    metadata = effect_dictionary()
    for effect in metadata["types"]:
        templates = [template for template in metadata["templates"] if template["effect_type"] == effect["code"]]
        for template in templates or [{"code": "", "label": ""}]:
            dictionary.append([effect["label"], effect["code"], "、".join(PARAMETER_LABELS[field] for field in effect["fields"]), template["code"], template["label"]])
    for column in ("A", "B", "C", "D", "E"):
        dictionary.column_dimensions[column].width = 28
    return workbook


def export_workbook(items):
    output = BytesIO()
    workbook = _workbook(items)
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _integer(value):
    if value is None or value == "":
        return None
    if type(value) is int:
        return value
    if type(value) is float and isfinite(value) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.isascii() and value.lstrip("-").isdigit():
        try:
            return int(value)
        except ValueError:
            pass
    return value


def parse_workbook(data, filename):
    if not filename or not filename.lower().endswith(".xlsx"):
        raise ValueError("仅支持 .xlsx；请在 Excel 中另存为 .xlsx 后重新上传")
    if not data or len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("文件为空或超过 5MB；请缩小文件后重新上传")
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 300 or sum(entry.file_size for entry in entries) > 30 * 1024 * 1024:
                raise ValueError("文件解压后过大；请使用模板，仅保留商品数据")
            for entry in entries:
                name = entry.filename.lower()
                if "vbaproject" in name or "externallinks/" in name or name.endswith(".bin"):
                    raise ValueError("不支持宏、外部链接或嵌入对象；请移除后重新上传")
                if name.endswith((".xml", ".rels")):
                    elements = iterparse(BytesIO(archive.read(entry)), events=("end",), forbid_dtd=True, forbid_entities=True, forbid_external=True)
                    for count, (_, element) in enumerate(elements, 1):
                        if count > 300000:
                            raise ValueError("工作簿结构过大；请使用模板，仅保留必要商品数据")
                        if name.startswith("xl/worksheets/"):
                            if element.tag.endswith("}row") and int(element.attrib.get("r", "1")) > MAX_ROWS + 1:
                                raise ValueError("最多导入 2000 行；请拆分文件，并删除多余格式行")
                            if element.tag.endswith("}sheetProtection"):
                                raise ValueError("请解除工作表保护后再上传")
                        element.clear()
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("文件损坏、不是有效的 Excel 或包含不安全内容；请使用模板重新保存为 .xlsx") from error
    try:
        if "商品列表" not in workbook.sheetnames:
            raise ValueError("缺少“商品列表”工作表；请下载最新模板填写")
        worksheet = workbook["商品列表"]
        worksheet.reset_dimensions()
        headers = next(worksheet.iter_rows(min_row=1, max_row=1, max_col=40, values_only=True))
        names = [value.strip() if isinstance(value, str) else value for value in headers]
        reverse = {FIELD_LABELS[field]: field for field in COLUMNS}
        missing = [FIELD_LABELS[field] for field in COLUMNS if field in REQUIRED and FIELD_LABELS[field] not in names]
        if missing:
            raise ValueError("缺少必需列：" + "、".join(missing) + "；请使用最新模板补齐列名")
        nonempty = [name for name in names if name is not None]
        if len(set(nonempty)) != len(nonempty):
            raise ValueError("列名重复；每个字段只保留一列")
        unknown = [name for name in nonempty if name not in reverse and name not in REPORT_COLUMNS]
        if unknown:
            raise ValueError("未知列：" + "、".join(str(name) for name in unknown) + "；请恢复模板列名")
        rows, errors, report_rows = [], [], []
        for line, cells in enumerate(worksheet.iter_rows(min_row=2, max_col=40), 2):
            if line > MAX_ROWS + 1:
                raise ValueError("最多导入 2000 行；请拆分文件")
            data_cells = [(index, cell) for index, cell in enumerate(cells) if names[index] in reverse]
            if not any(cell.value is not None and cell.value != "" for _, cell in data_cells):
                continue
            raw = {reverse[names[index]]: (
                cell.value.strip() if isinstance(cell.value, str)
                else cell.value if cell.value is None or isinstance(cell.value, (int, float, bool))
                else str(cell.value)
            ) for index, cell in data_cells}
            report_rows.append({"row": line, "raw": raw})
            for index, cell in data_cells:
                if cell.data_type in ("f", "e"):
                    errors.append(shop_issue(line, reverse[names[index]], str(cell.value), "公式或 Excel 错误值不能导入", "粘贴为实际文字或数字值，不使用公式"))
            row = {"row": line, "operation": raw.get("operation"), "public_number": _integer(raw.get("public_number"))}
            for field in READ_ONLY:
                if raw.get(field) is not None:
                    row[field] = _integer(raw[field]) if field in ("configuration_version", "base_stock") else raw[field]
            values = {field: raw.get(field) for field in (
                "name", "description", "price", "stock", "unlimited_stock", "enabled",
                "category", "daily_purchase_limit", "minimum_rank_order", "effect_type",
            )}
            for field in ("price", "stock", "daily_purchase_limit", "minimum_rank_order"):
                values[field] = _integer(values[field])
            for field in ("unlimited_stock", "enabled"):
                value = values[field]
                values[field] = {"是": True, "否": False}.get(value, value) if isinstance(value, str) else value
            code = next((code for code, label in EFFECT_LABELS.items() if values["effect_type"] == label), values["effect_type"])
            values["effect_type"] = None if code == "none" else code
            values["effect_config"] = {
                field: (raw.get(field) if field == "template" else _integer(raw.get(field)))
                for field in PARAMETER_LABELS if raw.get(field) is not None
            }
            row["values"] = values
            rows.append(row)
        return rows, errors, report_rows
    finally:
        workbook.close()


def error_report(report_rows, errors):
    workbook = _workbook([])
    worksheet = workbook["商品列表"]
    for offset, label in enumerate(REPORT_COLUMNS, len(COLUMNS) + 1):
        worksheet.cell(1, offset, label)
        worksheet.column_dimensions[get_column_letter(offset)].width = 65
    grouped = {}
    for error in errors:
        grouped.setdefault(error.get("row"), []).append(error)
    for entry in report_rows:
        line, raw = entry["row"], entry["raw"]
        for column, field in enumerate(COLUMNS, 1):
            _text_cell(worksheet.cell(line, column), raw.get(field))
        issues = grouped.get(line, [])
        values = ["错误" if issues else "通过", "；".join(f"{issue['column']}：{issue['message']}" for issue in issues), "；".join(issue["suggestion"] for issue in issues)]
        for offset, value in enumerate(values, len(COLUMNS) + 1):
            _text_cell(worksheet.cell(line, offset), value)
        for issue in issues:
            field = next((field for field in COLUMNS if FIELD_LABELS[field] == issue["column"]), None)
            if field:
                worksheet.cell(line, COLUMNS.index(field) + 1).fill = PatternFill("solid", fgColor="FFE1E1")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return b64encode(output.getvalue()).decode("ascii")
