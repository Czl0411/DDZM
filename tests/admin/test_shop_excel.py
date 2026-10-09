from base64 import b64decode
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from dzmm_bot.admin.app import create_app as create_admin_app
from dzmm_bot.admin.core_client import CoreClient
from dzmm_bot.admin.repository import AdminRepository
from dzmm_bot.admin.shop_excel import COLUMNS, export_workbook, parse_workbook
from dzmm_bot.core.app import create_app as create_core_app
from dzmm_bot.core.repository import CoreRepository
from dzmm_bot.core.schema import Base, ItemRecord, ShopChangeBatchRecord


NOW = datetime(2026, 10, 4, 4, tzinfo=UTC)
HEADERS = {"X-Admin-Token": "admin-secret"}


@pytest.fixture
def context(tmp_path):
    engine = create_engine("sqlite+pysqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    repository = CoreRepository(factory)
    repository.list_ranks()
    core_app = create_core_app(repository, "core-secret", clock=lambda: NOW)
    core_client = TestClient(core_app, headers={"X-Core-Token": "core-secret"})
    core = CoreClient("http://testserver", "core-secret", client=core_client)
    admin_repository = AdminRepository(factory)
    client = TestClient(create_admin_app("admin-secret", core, repository=admin_repository, profile_upload_dir=tmp_path / "profiles"))
    return client, repository, factory, admin_repository


def rewrite(data, changes):
    workbook = load_workbook(BytesIO(data))
    worksheet = workbook["商品列表"]
    for (line, field), value in changes.items():
        worksheet.cell(line, COLUMNS.index(field) + 1, value)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def upload(client, data, filename="products.xlsx", headers=HEADERS, sync=False):
    return client.post("/api/game/shop/excel/preview", headers=headers,
        files={"file": (filename, data)}, data={"synchronize_stock": str(sync).lower()})


def confirm(client, preview, headers=HEADERS):
    return client.post(f"/api/game/shop/changes/{preview['batch_id']}/confirm", headers=headers,
        json={"acknowledgements": [warning["code"] for warning in preview["warnings"]]})


def test_export_roundtrip_and_template_have_three_sheets(context):
    client, repository, _, _ = context
    response = client.get("/api/game/shop/excel/export", headers=HEADERS)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["商品列表", "填写说明", "效果字典"]
    assert workbook["商品列表"].max_row == 24
    workbook.close()
    preview = upload(client, response.content)
    assert preview.status_code == 200, preview.text
    assert preview.json()["summary"]["unchanged"] == 23
    assert not preview.json()["errors"]
    assert confirm(client, preview.json()).json()["count"] == 0
    template = client.get("/api/game/shop/excel/template", headers=HEADERS)
    rows, errors, _ = parse_workbook(template.content, "template.xlsx")
    assert rows == errors == []


def test_excel_updates_name_and_effect_only_after_confirmation(context):
    client, repository, _, _ = context
    exported = client.get("/api/game/shop/excel/export", headers=HEADERS).content
    edited = rewrite(exported, {(2, "name"): "新名称", (2, "reward"): 11})
    before = repository.get_shop_catalog()["items"][0]
    response = upload(client, edited)
    assert response.status_code == 200, response.text
    preview = response.json()
    assert not preview["errors"] and preview["warnings"]
    assert repository.get_shop_catalog()["items"][0] == before
    assert confirm(client, preview).status_code == 200
    current = repository.get_shop_catalog()["items"][0]
    assert current["name"] == "新名称" and current["effect_config"] == {"reward": 11}
    assert current["id"] == before["id"]
    assert confirm(client, preview).json()["count"] == 1


def test_upload_returns_all_errors_suggestions_and_reuploadable_marked_report(context):
    client, _, factory, _ = context
    exported = client.get("/api/game/shop/excel/export", headers=HEADERS).content
    edited = rewrite(exported, {(2, "price"): "十币", (3, "description"): "x" * 201, (4, "name"): "=1+1"})
    preview = upload(client, edited).json()
    assert preview["batch_id"] is None and preview["summary"]["errors"] >= 3
    assert {issue["row"] for issue in preview["errors"]} >= {2, 3, 4}
    assert all(issue["suggestion"] for issue in preview["errors"])
    report = load_workbook(BytesIO(b64decode(preview["error_report"])))
    worksheet = report["商品列表"]
    assert worksheet.cell(4, COLUMNS.index("name") + 1).data_type == "s"
    assert worksheet.cell(2, len(COLUMNS) + 1).value == "错误"
    assert worksheet.cell(2, COLUMNS.index("price") + 1).fill.fgColor.rgb.endswith("FFE1E1")
    worksheet.cell(2, COLUMNS.index("price") + 1, 3)
    worksheet.cell(3, COLUMNS.index("description") + 1, "修正描述")
    worksheet.cell(4, COLUMNS.index("name") + 1, "修正名称")
    output = BytesIO()
    report.save(output)
    report.close()
    corrected = upload(client, output.getvalue()).json()
    assert not corrected["errors"]
    with factory() as session:
        assert session.scalar(select(ItemRecord.price).where(ItemRecord.public_number == 1)) == 3
        assert len(list(session.scalars(select(ShopChangeBatchRecord)))) == 1


def test_archived_export_is_skipped_by_default(context):
    client, repository, _, _ = context
    item = repository.get_shop_catalog()["items"][0]
    preview = client.post("/api/game/shop/changes/preview", headers=HEADERS, json={"rows": [
        {"row": 2, "operation": "delete", "public_number": item["public_number"], "values": {}}
    ]}).json()
    assert confirm(client, preview).status_code == 200
    exported = client.get("/api/game/shop/excel/export?include_deleted=true", headers=HEADERS).content
    preview = upload(client, exported).json()
    assert not preview["errors"] and preview["summary"]["restore"] == 0
    assert confirm(client, preview).json()["count"] == 0
    assert repository.get_shop_catalog(True)["items"][0]["deleted_at"]


def test_regular_admin_cannot_change_effect_and_cannot_spoof_role(context):
    client, _, _, admin_repository = context
    account = admin_repository.create_account("operator", "test-password")
    headers = {"X-Admin-Session": admin_repository.create_session(account.id, datetime.now(UTC))}
    exported = client.get("/api/game/shop/excel/export", headers=headers).content
    preview = upload(client, rewrite(exported, {(2, "reward"): 99}), headers=headers).json()
    assert any("超级管理员" in issue["message"] for issue in preview["errors"])
    assert client.post("/api/game/shop/changes/preview", headers=headers,
        json={"rows": [], "super_admin": True}).status_code == 422
    unchanged = upload(client, exported, headers=headers).json()
    assert not unchanged["errors"]
    assert confirm(client, unchanged, headers=headers).status_code == 200


@pytest.mark.parametrize("method,path", [
    ("get", "/api/game/shop/catalog"), ("get", "/api/game/shop/excel/export"),
    ("get", "/api/game/shop/excel/template"), ("post", "/api/game/shop/excel/preview"),
    ("post", "/api/game/shop/changes/preview"),
])
def test_new_routes_require_authentication(context, method, path):
    client, _, _, _ = context
    response = getattr(client, method)(path)
    assert response.status_code == 401


@pytest.mark.parametrize("filename,data", [
    ("products.xls", b"not excel"), ("products.xlsx", b"not excel"),
    ("products.xlsx", b""), ("products.xlsx", b"x" * (5 * 1024 * 1024 + 1)),
], ids=["old-format", "corrupt", "empty", "oversized"])
def test_bad_uploads_have_file_level_suggestions(context, filename, data):
    client, _, _, _ = context
    response = upload(client, data, filename)
    assert response.status_code == 422
    assert "请" in response.json()["detail"]


def test_missing_header_and_unsafe_xml_are_rejected(context):
    client, _, _, _ = context
    workbook = load_workbook(BytesIO(export_workbook([])))
    workbook["商品列表"].cell(1, COLUMNS.index("name") + 1, "改错列名")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = upload(client, output.getvalue())
    assert "缺少必需列" in response.json()["detail"]
    malicious = BytesIO()
    with ZipFile(malicious, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", '<!DOCTYPE x [<!ENTITY danger "boom">]><x>&danger;</x>')
    assert upload(client, malicious.getvalue()).status_code == 422


def test_transport_limit_rejects_large_body_before_parsing(context):
    client, _, _, _ = context
    response = upload(client, b"x" * (6 * 1024 * 1024))
    assert response.status_code == 413
    assert "5MB" in response.json()["detail"]


def test_transport_limit_also_checks_chunked_uploads(context):
    client, _, _, _ = context
    def chunks():
        yield b'--shop-boundary\r\nContent-Disposition: form-data; name="file"; filename="products.xlsx"\r\nContent-Type: application/octet-stream\r\n\r\n'
        for _ in range(6):
            yield b"x" * (1024 * 1024)
        yield b"\r\n--shop-boundary--\r\n"
    response = client.post("/api/game/shop/excel/preview", content=chunks(),
        headers={**HEADERS, "Content-Type": "multipart/form-data; boundary=shop-boundary"})
    assert response.status_code == 413
    assert "5MB" in response.json()["detail"]


def test_conflict_is_actionable_and_not_double_json_encoded(context):
    client, repository, factory, _ = context
    exported = client.get("/api/game/shop/excel/export", headers=HEADERS).content
    preview = upload(client, rewrite(exported, {(2, "name"): "准备改名"})).json()
    with factory.begin() as session:
        session.scalar(select(ItemRecord).where(ItemRecord.public_number == 1)).configuration_version += 1
    response = confirm(client, preview)
    assert response.status_code == 409
    assert response.json()["detail"].startswith("商品 #1")
    assert "重新预览" in response.json()["detail"]
    assert repository.get_shop_catalog()["items"][0]["name"] != "准备改名"


def test_export_does_not_turn_untrusted_text_into_formulas():
    rows = [{"public_number": 1, "name": "=HYPERLINK(\"https://invalid\")", "description": "@SUM(1)",
             "price": 1, "stock": 1, "effect_type": None, "effect_config": {}, "enabled": True, "unlimited_stock": False}]
    workbook = load_workbook(BytesIO(export_workbook(rows)))
    assert workbook["商品列表"].cell(2, COLUMNS.index("name") + 1).data_type == "s"
    workbook.close()


def test_shop_ui_has_import_preview_controls_and_safe_rendering():
    root = Path(__file__).resolve().parents[2] / "src/dzmm_bot/admin"
    html = (root / "templates/index.html").read_text(encoding="utf-8")
    script = (root / "static/admin.js").read_text(encoding="utf-8")
    for identifier in ("shop-excel-export", "shop-excel-template", "shop-excel-import", "shop-import-preview", "shop-sync-stock", "shop-include-deleted"):
        assert f'id="{identifier}"' in html
    assert "escapeHtml(issue.suggestion)" in script
    assert "data-shop-risk" in script
    assert "shopPreviewGeneration" in script
